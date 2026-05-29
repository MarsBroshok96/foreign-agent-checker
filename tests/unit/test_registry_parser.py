from datetime import date, datetime

import pytest
from openpyxl import Workbook

from fa_checker.domain.enums import EntityType, MatchType
from fa_checker.domain.models import Article
from fa_checker.matching.exact import find_exact_matches
from fa_checker.registry.parser import parse_registry_xlsx

SOURCE_URL = "https://minjust.gov.ru/registry"


def save_workbook(workbook: Workbook, path) -> None:
    workbook.save(path)


def test_parse_registry_xlsx_simple_minimal_format(tmp_path) -> None:
    path = tmp_path / "simple.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type", "aliases"])
    sheet.append(["1", "Варламов Илья Александрович", "person", "Илья Варламов; Варламов"])
    sheet.append(["2", "ООО Ромашка", "organization", "Ромашка"])
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL, snapshot_date=date(2026, 5, 22))

    assert len(entries) == 2
    assert entries[0].registry_id == "1"
    assert entries[0].full_name == "Варламов Илья Александрович"
    assert entries[0].entity_type == EntityType.PERSON
    assert entries[0].normalized_name == "варламов илья александрович"
    assert entries[0].aliases == ["Илья Варламов", "Варламов"]
    assert entries[0].registry_source_url == SOURCE_URL
    assert entries[0].registry_snapshot_date == date(2026, 5, 22)
    assert entries[1].entity_type == EntityType.ORGANIZATION
    assert entries[1].aliases == ["Ромашка"]


def test_parse_registry_xlsx_official_like_layout(tmp_path) -> None:
    path = tmp_path / "official.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["Реестр иностранных агентов"])
    sheet.append([datetime(2026, 5, 22)])
    sheet.append(
        [
            "№ п/п",
            "Полное наименование (прежнее наименование (в случае его изменения)) / "
            "ФИО «Псевдоним» (при наличии) (прежние ФИО (в случае их изменения))",
            "Основания для включения",
            "Дата принятия решения о включении",
            "Дата исключения",
            "Доменное имя информационного ресурса",
            "Тип иностранного агента",
            "Регистрационный номер",
            "Полное наименование или ФИО участников",
            "Дата опубликования решения о включении",
        ]
    )
    sheet.append(
        [
            "1208",
            "Проект «После»",
            "основание",
            "2025-01-01",
            None,
            "https://posle.media/ https://t.me/poslemedia",
            "Иные объединения лиц",
            "reg-1208",
            "Участник Один",
            "2025-01-02",
        ]
    )
    sheet.append(
        [
            "1207",
            "Чойнова Вилюя Борисовна",
            "основание",
            "2025-01-01",
            None,
            None,
            "Физические лица",
            "reg-1207",
            None,
            "2025-01-02",
        ]
    )
    sheet.append(
        [
            "560",
            "Варламов Илья Александрович",
            "основание",
            "2023-01-01",
            None,
            "https://varlamov.ru/",
            "Физические лица",
            "reg-560",
            None,
            "2023-01-02",
        ]
    )
    sheet.append(
        [
            "728",
            '"Телеканал Дождь"',
            "основание",
            "2024-01-01",
            None,
            "https://tvrain.tv",
            "Иные объединения лиц",
            "reg-728",
            None,
            "2024-01-02",
        ]
    )
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert len(entries) == 4
    by_id = {entry.registry_id: entry for entry in entries}
    assert by_id["1208"].aliases == ["После"]
    assert by_id["1208"].raw_fields["quoted_aliases"] == ["После"]
    assert by_id["1208"].entity_type == EntityType.PROJECT
    assert by_id["1208"].registry_snapshot_date == date(2026, 5, 22)
    assert by_id["1208"].raw_fields["resource_urls"] == [
        "https://posle.media/",
        "https://t.me/poslemedia",
    ]
    assert by_id["1208"].raw_fields["resource_domains"] == ["posle.media", "t.me"]
    assert by_id["1208"].raw_fields["participants"] == "Участник Один"
    assert "Участник Один" not in by_id["1208"].aliases
    assert by_id["560"].entity_type == EntityType.PERSON
    assert by_id["560"].raw_fields["resource_domains"] == ["varlamov.ru"]
    assert by_id["728"].aliases == ["Телеканал Дождь"]
    assert by_id["728"].raw_fields["quoted_aliases"] == ["Телеканал Дождь"]


def test_parse_registry_xlsx_supports_russian_simple_headers(tmp_path) -> None:
    path = tmp_path / "russian_headers.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["№", "ФИО", "тип", "алиасы"])
    sheet.append(["1", "Варламов Илья Александрович", "физлицо", "Илья Варламов"])
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert len(entries) == 1
    assert entries[0].registry_id == "1"
    assert entries[0].entity_type == EntityType.PERSON
    assert entries[0].aliases == ["Илья Варламов"]


def test_parse_registry_xlsx_skips_empty_full_name_rows(tmp_path) -> None:
    path = tmp_path / "empty_rows.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type"])
    sheet.append(["1", None, "person"])
    sheet.append(["2", "ООО Ромашка", "organization"])
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert len(entries) == 1
    assert entries[0].full_name == "ООО Ромашка"


def test_parse_registry_xlsx_missing_full_name_column_raises(tmp_path) -> None:
    path = tmp_path / "missing_full_name.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "entity_type", "aliases"])
    sheet.append(["1", "person", "alias"])
    save_workbook(workbook, path)

    with pytest.raises(ValueError, match="missing a full_name"):
        parse_registry_xlsx(path, SOURCE_URL)


def test_parse_registry_xlsx_unknown_entity_type_maps_to_unknown(tmp_path) -> None:
    path = tmp_path / "unknown_type.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type"])
    sheet.append(["1", "Неизвестная сущность", "непонятный тип"])
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert entries[0].entity_type == EntityType.UNKNOWN


def test_parse_registry_xlsx_empty_aliases_produce_empty_list(tmp_path) -> None:
    path = tmp_path / "empty_aliases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type", "aliases"])
    sheet.append(["1", "ООО Ромашка", "organization", "  "])
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert entries[0].aliases == []


def test_parse_registry_xlsx_deduplicates_quoted_aliases_with_explicit_aliases(tmp_path) -> None:
    path = tmp_path / "dedupe_aliases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type", "aliases"])
    sheet.append(["1", "Проект «После»", "project", "После; Другое"])
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert entries[0].aliases == ["После", "Другое"]
    assert entries[0].raw_fields["explicit_aliases"] == ["После", "Другое"]
    assert entries[0].raw_fields["quoted_aliases"] == ["После"]


def test_parse_registry_xlsx_stores_explicit_aliases_traceably(tmp_path) -> None:
    path = tmp_path / "explicit_aliases.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type", "aliases"])
    sheet.append(
        ["1", "Варламов Илья Александрович", "person", "Илья Варламов; Варламов"]
    )
    save_workbook(workbook, path)

    entries = parse_registry_xlsx(path, SOURCE_URL)

    assert entries[0].aliases == ["Илья Варламов", "Варламов"]
    assert entries[0].raw_fields["explicit_aliases"] == ["Илья Варламов", "Варламов"]


def test_parse_registry_xlsx_entries_feed_exact_matcher(tmp_path) -> None:
    path = tmp_path / "matcher.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["registry_id", "full_name", "entity_type"])
    sheet.append(["560", "Варламов Илья Александрович", "Физические лица"])
    save_workbook(workbook, path)
    entries = parse_registry_xlsx(path, SOURCE_URL)
    article = Article(
        url="https://www.rambler.ru/example",
        source_domain="www.rambler.ru",
        text="Илья Варламов прокомментировал ситуацию.",
    )

    matches = find_exact_matches(article, entries)

    assert len(matches) == 1
    assert matches[0].match_type == MatchType.EXACT
    assert matches[0].requires_disambiguation is False
