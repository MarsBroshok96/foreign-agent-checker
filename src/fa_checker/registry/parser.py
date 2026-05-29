"""Local XLSX parser for Ministry of Justice registry snapshots."""

import re
from datetime import date, datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from fa_checker.article.normalizer import normalize_for_matching
from fa_checker.domain.enums import EntityType
from fa_checker.domain.models import RegistryEntry

HEADER_SCAN_ROWS = 20
CANONICAL_FIELDS = {
    "registry_id",
    "full_name",
    "entity_type",
    "aliases",
    "resource_urls",
    "participants",
    "included_at",
    "excluded_at",
    "published_at",
    "registration_number",
}
RESOURCE_SPLIT_RE = re.compile(r"[\s;,]+")
ANGLE_QUOTE_RE = re.compile(r"«([^»]+)»")
DOUBLE_QUOTE_RE = re.compile(r'"([^"]+)"')
DATE_RE = re.compile(r"(\d{4}-\d{2}-\d{2}|\d{2}\.\d{2}\.\d{4})")


class RegistryParser:
    """Scaffold class kept until parser orchestration is designed."""

    def parse(self, raw_content: bytes) -> list[RegistryEntry]:
        raise NotImplementedError("Registry parsing is not implemented in the scaffold phase.")


def parse_registry_xlsx(
    path: str | Path,
    registry_source_url: str,
    snapshot_date: date | None = None,
) -> list[RegistryEntry]:
    """Parse a local XLSX registry snapshot into registry entries."""
    workbook = load_workbook(path, data_only=True)
    sheet = workbook.active
    header_row = _detect_header_row(sheet)
    columns = _build_column_map(sheet[header_row])

    if "full_name" not in columns:
        msg = "Registry XLSX is missing a full_name/name/official full-name column."
        raise ValueError(msg)

    resolved_snapshot_date = snapshot_date or _infer_snapshot_date(sheet)
    entries: list[RegistryEntry] = []

    for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        raw_row = _read_raw_row(row, columns)
        full_name = _string_or_empty(raw_row.get("full_name"))
        if not full_name:
            continue

        raw_registry_id = raw_row.get("registry_id")
        registry_id = _string_or_none(raw_registry_id)
        resource_urls = _split_resource_urls(raw_row.get("resource_urls"))
        explicit_aliases = _split_aliases(raw_row.get("aliases"))
        quoted_aliases = _extract_quoted_aliases(full_name)
        aliases = _collect_aliases(explicit_aliases, quoted_aliases)
        raw_fields = _build_raw_fields(
            raw_row,
            resource_urls,
            explicit_aliases,
            quoted_aliases,
        )

        entries.append(
            RegistryEntry(
                registry_id=registry_id,
                full_name=full_name,
                entity_type=_map_entity_type(raw_row.get("entity_type")),
                normalized_name=normalize_for_matching(full_name),
                aliases=aliases,
                registry_source_url=registry_source_url,
                registry_snapshot_date=resolved_snapshot_date,
                raw_fields=raw_fields,
            )
        )

    return entries


def _detect_header_row(sheet: Any) -> int:
    for row_index in range(1, min(sheet.max_row, HEADER_SCAN_ROWS) + 1):
        canonical_headers = {
            canonical
            for cell in sheet[row_index]
            if (canonical := _canonical_header(cell.value)) is not None
        }
        if len(canonical_headers) >= 2:
            return row_index
    msg = "Could not detect registry table header row in the first 20 rows."
    raise ValueError(msg)


def _build_column_map(header_cells: Any) -> dict[str, int]:
    columns: dict[str, int] = {}
    for index, cell in enumerate(header_cells):
        canonical = _canonical_header(cell.value)
        if canonical is not None and canonical not in columns:
            columns[canonical] = index
    return columns


def _canonical_header(value: Any) -> str | None:
    normalized = normalize_for_matching(_string_or_empty(value))
    if not normalized:
        return None

    if normalized in {"registry_id", "id", "№", "№ п/п", "номер", "номер п/п"}:
        return "registry_id"
    if normalized in {"registration_number", "регистрационный номер"}:
        return "registration_number"
    if normalized in {"participants", "участники", "полное наименование или фио участников"}:
        return "participants"
    if "участников" in normalized and ("фио" in normalized or "наименование" in normalized):
        return "participants"
    if normalized in {"full_name", "name", "фио", "полное наименование"}:
        return "full_name"
    if (
        "полное наименование" in normalized
        and "фио" in normalized
        and "участников" not in normalized
    ):
        return "full_name"
    if normalized in {"entity_type", "type", "тип", "тип иностранного агента"}:
        return "entity_type"
    if normalized in {"aliases", "alias", "алиасы", "алиас", "псевдоним"}:
        return "aliases"
    if normalized in {"resource_urls", "resource_url", "доменное имя информационного ресурса"}:
        return "resource_urls"
    if "доменное имя" in normalized and "информационного ресурса" in normalized:
        return "resource_urls"
    if normalized in {"included_at", "дата принятия решения о включении"}:
        return "included_at"
    if normalized in {"excluded_at", "дата исключения"}:
        return "excluded_at"
    if normalized in {"published_at", "дата опубликования решения о включении"}:
        return "published_at"
    return None


def _infer_snapshot_date(sheet: Any) -> date | None:
    if sheet.max_row < 2:
        return None
    for cell in sheet[2]:
        parsed = _parse_date(cell.value)
        if parsed is not None:
            return parsed
    return None


def _parse_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, int | float):
        try:
            return from_excel(value).date()
        except (TypeError, ValueError, OverflowError):
            return None
    if isinstance(value, str):
        match = DATE_RE.search(value.strip())
        if match is None:
            return None
        date_text = match.group(1)
        for date_format in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(date_text, date_format).date()
            except ValueError:
                continue
    return None


def _read_raw_row(row: tuple[Any, ...], columns: dict[str, int]) -> dict[str, Any]:
    return {
        canonical: row[index] if index < len(row) else None
        for canonical, index in columns.items()
        if canonical in CANONICAL_FIELDS
    }


def _build_raw_fields(
    raw_row: dict[str, Any],
    resource_urls: list[str],
    explicit_aliases: list[str],
    quoted_aliases: list[str],
) -> dict[str, Any]:
    raw_fields = {field: raw_row.get(field) for field in CANONICAL_FIELDS}
    raw_fields["resource_urls"] = resource_urls
    raw_fields["resource_domains"] = _extract_resource_domains(resource_urls)
    raw_fields["participants"] = raw_row.get("participants")
    raw_fields["explicit_aliases"] = explicit_aliases
    raw_fields["quoted_aliases"] = quoted_aliases
    raw_fields["raw_row"] = dict(raw_row)
    return raw_fields


def _collect_aliases(explicit_aliases: list[str], quoted_aliases: list[str]) -> list[str]:
    aliases: list[str] = []
    for alias in explicit_aliases:
        _add_unique(aliases, alias)
    for alias in quoted_aliases:
        _add_unique(aliases, alias)
    return aliases


def _split_aliases(raw_aliases: Any) -> list[str]:
    if raw_aliases is None:
        return []
    return [
        alias.strip()
        for alias in str(raw_aliases).split(";")
        if alias.strip()
    ]


def _extract_quoted_aliases(text: str) -> list[str]:
    aliases: list[str] = []
    for pattern in (ANGLE_QUOTE_RE, DOUBLE_QUOTE_RE):
        for match in pattern.finditer(text):
            _add_unique(aliases, match.group(1).strip())
    return aliases


def _add_unique(items: list[str], value: str) -> None:
    if value and value not in items:
        items.append(value)


def _split_resource_urls(value: Any) -> list[str]:
    if value is None:
        return []
    urls: list[str] = []
    for item in RESOURCE_SPLIT_RE.split(str(value)):
        cleaned = item.strip().strip(".,;")
        if cleaned:
            _add_unique(urls, cleaned)
    return urls


def _extract_resource_domains(resource_urls: list[str]) -> list[str]:
    domains: list[str] = []
    for resource_url in resource_urls:
        parsed = urlparse(resource_url)
        if parsed.scheme not in {"http", "https"} or not parsed.hostname:
            continue
        _add_unique(domains, parsed.hostname.rstrip(".").lower())
    return domains


def _map_entity_type(value: Any) -> EntityType:
    normalized = normalize_for_matching(_string_or_empty(value))
    if not normalized:
        return EntityType.UNKNOWN
    if normalized in {
        "person",
        "physical_person",
        "физическое лицо",
        "физические лица",
        "физлицо",
        "человек",
        "персона",
    }:
        return EntityType.PERSON
    if normalized in {
        "organization",
        "org",
        "юрлицо",
        "юридическое лицо",
        "юридические лица",
        "организация",
        "иностранные структуры без образования юридического лица",
        "общественные объединения, действующие без образования юридического лица",
    }:
        return EntityType.ORGANIZATION
    if normalized in {"media", "сми", "медиа"}:
        return EntityType.MEDIA
    if normalized in {"project", "проект", "иные объединения лиц"}:
        return EntityType.PROJECT
    return EntityType.UNKNOWN


def _string_or_empty(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _string_or_none(value: Any) -> str | None:
    text = _string_or_empty(value)
    return text or None
